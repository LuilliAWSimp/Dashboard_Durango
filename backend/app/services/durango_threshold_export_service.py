"""Exportaciones individuales H02 usando exclusivamente el historico del backend."""
from __future__ import annotations

from html import escape
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle

from app.services.durango_flow_threshold import FlowThreshold, evaluate, number, summarize
from app.services.water_module_history_pdf_service import _chart_drawing

FLAGGED = {'Bajo umbral', 'Sobre umbral', 'Fuera de rango'}
HEADERS = ['Inicio', 'Fin', 'Prom. (L/s)', 'Min. (L/s)', 'Max. (L/s)',
           'Volumen (m³)', 'Tot. cierre (m³)', 'Muestras', 'Estado umbral']


def _rows(history: dict[str, Any], rule: FlowThreshold) -> list[tuple[Any, ...]]:
    output = []
    for point in history.get('points') or []:
        if point.get('data_status') == 'future_interval':
            continue
        output.append((
            point.get('bucket_start'), point.get('bucket_end'),
            number(point.get('flow_avg_lps')), number(point.get('flow_min_lps')),
            number(point.get('flow_max_lps')),
            number(point.get('volume_m3')) if point.get('volume_reliable') else None,
            number(point.get('totalizer_close_m3')),
            int(point.get('samples') or 0), evaluate(point, rule),
        ))
    return output


def _summary_text(counts: dict[str, int]) -> str:
    return (f"{counts['breached']} intervalos fuera de rango de {counts['total']} examinados; "
            f"{counts['below']} bajo mínimo, {counts['above']} sobre máximo; "
            f"{counts['within']} dentro de rango y {counts['no_reading']} sin lectura. "
            'Los conteos bajo/sobre pueden coincidir en un mismo intervalo.')


def build_threshold_pdf(history: dict[str, Any], rule: FlowThreshold) -> tuple[bytes, str]:
    rows = _rows(history, rule)
    counts = summarize([p for p in history.get('points') or [] if p.get('data_status') != 'future_interval'], rule)
    name = str(history.get('name') or 'Sensor')
    sensor = str(history.get('sensor_id') or history.get('operational_key') or '')
    start = str(history.get('start_date') or '')
    end = str(history.get('end_date') or '')
    aggregation = str(history.get('aggregation') or '')
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), leftMargin=10*mm, rightMargin=10*mm,
                            topMargin=10*mm, bottomMargin=12*mm, title=f'Umbrales de flujo - {name}',
                            author='Dashboard ARCA - Durango')
    base = getSampleStyleSheet()
    title_style = ParagraphStyle('H02Title', parent=base['Heading1'], fontSize=17, textColor=colors.HexColor('#163344'), spaceAfter=6)
    text_style = ParagraphStyle('H02Text', parent=base['Normal'], fontSize=8, leading=11, textColor=colors.HexColor('#334155'), spaceAfter=6)
    small_style = ParagraphStyle('H02Cell', parent=base['Normal'], fontSize=6.8, leading=8, textColor=colors.HexColor('#163344'))
    head_style = ParagraphStyle('H02Head', parent=small_style, fontName='Helvetica-Bold', textColor=colors.white)
    story: list[Any] = [
        Paragraph('DASHBOARD ARCA · PLANTA DURANGO', text_style),
        Paragraph(f'Revisión histórica de umbrales · {escape(name)}', title_style),
        Paragraph(f'Sensor: {escape(sensor)} · Periodo: {escape(start)} a {escape(end)} · Agrupación: {escape(aggregation)}', text_style),
        Paragraph(escape(rule.label()), text_style),
        Paragraph(escape(_summary_text(counts)), text_style),
    ]
    # La grafica PDF comparte los buckets de este mismo historico; umbrales dibujados en L/s.
    graph_rows = [{'label': str(row[0])[5:16], 'flow': row[2]} for row in rows]
    graph_series = [{'key': 'flow', 'name': 'Flujo promedio', 'metric': 'flow', 'axis': 'left', 'chart_type': 'line', 'color': '#096D94'}]
    if rule.direction in {'below', 'both'}:
        for row in graph_rows:
            row['minimum'] = rule.min_value
        graph_series.append({'key': 'minimum', 'name': 'Limite minimo', 'metric': 'threshold', 'axis': 'left', 'chart_type': 'line', 'color': '#F59E0B'})
    if rule.direction in {'above', 'both'}:
        for row in graph_rows:
            row['maximum'] = rule.max_value
        graph_series.append({'key': 'maximum', 'name': 'Limite maximo', 'metric': 'threshold', 'axis': 'left', 'chart_type': 'line', 'color': '#FB7185'})
    if graph_rows:
        # Limitar solo la resolucion del dibujo; la tabla y el diagnostico incluyen TODOS los buckets.
        stride = max(1, (len(graph_rows) + 299)//300)
        story.append(_chart_drawing(graph_rows[::stride], graph_series, 'Flujo (L/s)', ''))
        story.append(Spacer(1, 3*mm))
    story.append(Paragraph('Cada incidencia corresponde al intervalo de agrupación, no al segundo exacto del evento. Un cero medido es válido; una lectura ausente se deja vacía.', text_style))
    table_data = [[Paragraph(escape(h), head_style) for h in HEADERS]]
    for row in rows:
        values = []
        for idx, val in enumerate(row):
            displayed = '' if val is None else f'{val:,.2f}' if isinstance(val, float) else str(val)
            values.append(Paragraph(escape(displayed), small_style))
        table_data.append(values)
    if rows:
        table = LongTable(table_data, repeatRows=1, hAlign='LEFT',
                          colWidths=[32*mm, 32*mm, 21*mm, 21*mm, 21*mm, 25*mm, 26*mm, 16*mm, 52*mm])
        styles = [('BACKGROUND', (0,0),(-1,0), colors.HexColor('#163344')),
                  ('ROWBACKGROUNDS', (0,1),(-1,-1), [colors.white, colors.HexColor('#F2F8FC')]),
                  ('GRID',(0,0),(-1,-1),0.25,colors.HexColor('#CFDEE6')),
                  ('VALIGN',(0,0),(-1,-1),'MIDDLE'),
                  ('BOTTOMPADDING',(0,0),(-1,-1),3), ('TOPPADDING',(0,0),(-1,-1),3)]
        for idx, row in enumerate(rows, start=1):
            if row[-1] in FLAGGED:
                styles.append(('BACKGROUND', (0,idx),(-1,idx),colors.HexColor('#FFF3C4')))
        table.setStyle(TableStyle(styles))
        story.append(table)
    doc.build(story)
    return buffer.getvalue(), f'umbral-flujo-durango-{sensor}-{start}-{end}.pdf'


def build_threshold_excel(history: dict[str, Any], rule: FlowThreshold) -> tuple[bytes, str]:
    rows = _rows(history, rule)
    points = [p for p in history.get('points') or [] if p.get('data_status') != 'future_interval']
    counts = summarize(points, rule)
    wb = Workbook()
    ws = wb.active
    ws.title = 'Umbrales de flujo'
    ws.append(['PLANTA DURANGO · REVISION HISTORICA DE UMBRALES'])
    ws.append(['Elemento', history.get('name') or 'Sensor'])
    ws.append(['Periodo', f"{history.get('start_date')} a {history.get('end_date')}"])
    ws.append(['Agrupación', history.get('aggregation')])
    ws.append(['Criterio', rule.label()])
    ws.append(['Diagnóstico', _summary_text(counts)])
    ws.append([])
    ws.append(HEADERS)
    ws.merge_cells('A1:I1')
    for row in rows:
        ws.append(list(row))
        excel_idx = ws.max_row
        if row[-1] in FLAGGED:
            for cell in ws[excel_idx]:
                cell.fill = PatternFill('solid', fgColor='FFF3C4')
        for cell in ws[excel_idx][2:7]:
            cell.number_format = '#,##0.00'
    for cell in ws[1]:
        cell.fill = PatternFill('solid', fgColor='163344')
        cell.font = Font(color='FFFFFF', bold=True, size=14)
    for cell in ws[8]:
        cell.fill = PatternFill('solid', fgColor='096D94')
        cell.font = Font(color='FFFFFF', bold=True)
        cell.alignment = Alignment(wrap_text=True)
    for rownum in range(2, 7):
        ws.cell(rownum, 1).font = Font(bold=True, color='163344')
    for index, width in enumerate([23, 23, 17, 17, 17, 20, 23, 13, 23], start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = 'C9'
    ws.auto_filter.ref = f'A8:I{ws.max_row}'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.print_title_rows = '1:8'
    out = BytesIO()
    wb.save(out)
    sensor = str(history.get('sensor_id') or history.get('operational_key') or 'sensor')
    return out.getvalue(), f'umbral-flujo-durango-{sensor}-{history.get("start_date")}-{history.get("end_date")}.xlsx'
