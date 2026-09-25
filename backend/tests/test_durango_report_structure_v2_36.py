from io import BytesIO
from unittest.mock import patch

from openpyxl import load_workbook

from app.services.water_daily_report_service import build_daily_water_report_excel, get_daily_water_report


def item(name: str, module: str, key: str, volume: float | None, activity: str = 'Con actividad') -> dict:
    return {
        'name': name,
        'module': module,
        'operational_key': key,
        'current_flow': 1.5,
        'flow_unit': 'L/s',
        'period_open_m3': 100.0,
        'period_close_m3': 101.0,
        'period_m3': volume,
        'period_m3_reliable': volume is not None,
        'validated_volume_m3': volume,
        'discarded_volume_m3': 0.0,
        'activity': activity,
        'communication': 'Actualizado',
        'last_update': '2026-09-25T12:00:00',
        'samples': 10,
    }


def review_payload() -> dict:
    wells = [item('Pozo 1', 'well', 'pozo_1', 10.0)]
    lines = [item('Línea 1', 'line', 'linea_1', 8.0)]
    flows = [
        item('Lavadora Línea 2', 'flow', 'lavadora_linea_2', 4.0),
        item('Jarabes', 'flow', 'jarabes', 2.0),
    ]
    return {
        'source_status': 'operational',
        'validated_segment_start': '2026-08-04T18:16:00',
        'crosses_scada_cutover': False,
        'legacy_notice': None,
        'modules': {
            'wells': {'items': wells, 'summary': {'total_m3': 10.0, 'active_count': 1, 'review_count': 0}},
            'lines': {'items': lines, 'summary': {'total_m3': 8.0, 'active_count': 1, 'review_count': 0}},
            'flows': {'items': flows, 'summary': {'total_m3': 6.0, 'active_count': 2, 'review_count': 0}},
        },
        'shifts': {'shifts': []},
    }


def build_report() -> dict:
    with patch('app.services.water_daily_report_service.get_daily_water_review', return_value=review_payload()), patch(
        'app.services.water_daily_report_service.get_water_history_module',
        return_value={'series': [], 'aggregation': 'quarter_hour'},
    ):
        return get_daily_water_report('2026-09-25')


def test_presentation_contract_has_canonical_summary_and_module_order() -> None:
    report = build_report()
    presentation = report['presentation']
    assert [card['key'] for card in presentation['summary_cards']] == [
        'wells_volume', 'lines_volume', 'washers_volume', 'jarabes_volume', 'active_items', 'attention_items'
    ]
    assert [section['label'] for section in presentation['sections']] == ['Pozos', 'Líneas', 'Lavadoras', 'Jarabes']
    assert [column['label'] for column in presentation['table_columns']] == [
        'Elemento', 'Flujo actual', 'Totalizador inicial', 'Totalizador final', 'Volumen del periodo',
        'Actividad', 'Estado de datos', 'Comunicación', 'Última lectura',
    ]
    assert all('operativo' not in str(card['label']).lower() for card in presentation['summary_cards'])


def test_excel_uses_same_summary_without_cross_module_total() -> None:
    content, _ = build_daily_water_report_excel(build_report())
    workbook = load_workbook(BytesIO(content), data_only=False)
    summary = workbook['Resumen']
    labels = [summary.cell(row, 1).value for row in range(2, summary.max_row + 1)]
    assert labels[3:9] == [
        'Volumen bombeado de pozos (m³)',
        'Volumen consumido de líneas (m³)',
        'Volumen consumido de lavadoras (m³)',
        'Volumen consumido de Jarabes (m³)',
        'Con actividad',
        'Con atención',
    ]
    assert not any(label and 'Total validado operativo' in str(label) for label in labels)
    assert not any(label and 'Subtotal validado operativo' in str(label) for label in labels)
    assert workbook.sheetnames[:6] == ['Resumen', 'Pozos', 'Líneas', 'Lavadoras', 'Jarabes', 'Turnos']
    assert workbook['Turnos'].max_column == 7
    assert workbook['Turnos']['G1'].value == 'Estado'
