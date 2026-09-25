from io import BytesIO
from unittest.mock import patch

from openpyxl import load_workbook

from app.services.water_daily_report_service import build_daily_water_report_excel, get_daily_water_report


def _item(name: str, module: str, key: str, volume: float | None, samples: int = 10) -> dict:
    return {
        'name': name,
        'module': module,
        'operational_key': key,
        'current_flow': 1.5,
        'flow_unit': 'L/s',
        'period_open_m3': 100.0,
        'period_close_m3': 101.0,
        'period_m3': volume,
        'period_m3_reliable': volume is not None,
        'validated_volume_m3': volume,
        'activity': 'Con actividad' if samples else 'Sin registros',
        'communication': 'Actualizado',
        'last_update': '2026-09-25T12:00:00' if samples else None,
        'samples': samples,
        'quality_status': 'validated' if volume is not None else ('review' if samples else 'no_data'),
        'quality_label': 'Validado' if volume is not None else ('Dato en revisión' if samples else 'Sin datos'),
    }


def _review() -> dict:
    wells = [_item('Pozo 1', 'well', 'pozo_1', 10.0)]
    lines = [_item('Línea 1', 'line', 'linea_1', None, 10)]
    flows = [_item('Lavadora Línea 2', 'flow', 'lavadora_linea_2', 4.0), _item('Jarabes', 'flow', 'jarabes', None, 0)]
    return {
        'source_status': 'operational',
        'validated_segment_start': '2026-08-04T18:16:00',
        'crosses_scada_cutover': False,
        'legacy_notice': None,
        'modules': {
            'wells': {'items': wells, 'summary': {'total_m3': 10.0, 'active_count': 1, 'review_count': 0}},
            'lines': {'items': lines, 'summary': {'total_m3': None, 'active_count': 1, 'review_count': 1}},
            'flows': {'items': flows, 'summary': {'total_m3': 4.0, 'active_count': 1, 'review_count': 1}},
        },
        'shifts': {'shifts': []},
    }


def _history(module: str, start_date: str, end_date: str, aggregation: str) -> dict:
    return {
        'module': module,
        'aggregation': aggregation,
        'series': [{
            'name': 'Elemento',
            'points': [{
                'bucket_start': f'{start_date}T00:00:00',
                'bucket_end': f'{start_date}T00:15:00',
                'aggregation': aggregation,
                'flow_avg_lps': 1.0,
                'samples': 10,
                'samples_received': 10,
                'samples_expected': 15,
                'coverage_percent': 66.7,
                'totalizer_open_m3': 1.0,
                'totalizer_close_m3': 2.0,
                'validated_volume_m3': None,
                'interval_state': 'Actividad parcial',
                'data_status': 'invalid_totalizer',
            }],
        }],
    }


def _report() -> dict:
    with patch('app.services.water_daily_report_service.get_daily_water_review', return_value=_review()), patch(
        'app.services.water_daily_report_service.get_water_history_module', side_effect=_history
    ):
        return get_daily_water_report('2026-09-25')


def test_excel_uses_operational_copy_in_summary_and_module_tables() -> None:
    content, _ = build_daily_water_report_excel(_report())
    wb = load_workbook(BytesIO(content), data_only=False)
    summary = {wb['Resumen'].cell(row, 1).value: wb['Resumen'].cell(row, 2).value for row in range(2, wb['Resumen'].max_row + 1)}
    assert 'Estado de datos' in summary
    assert 'Cobertura del reporte' not in summary
    assert 'Criterio de cálculo' not in summary
    assert wb['Pozos']['G2'].value == 'Datos completos'
    assert wb['Líneas']['G2'].value == 'Datos parciales'
    assert wb['Jarabes']['G2'].value == 'Sin datos'


def test_historical_excel_translates_internal_status_codes() -> None:
    content, _ = build_daily_water_report_excel(_report())
    wb = load_workbook(BytesIO(content), data_only=False)
    history = wb['Histórico Pozos']
    assert history['L1'].value == 'Disponibilidad (%)'
    assert history['O1'].value == 'Volumen del intervalo (m³)'
    assert history['P1'].value == 'Actividad del intervalo'
    assert history['Q1'].value == 'Estado de datos'
    assert history['Q2'].value == 'Datos parciales'
    assert history['Q2'].value != 'invalid_totalizer'
