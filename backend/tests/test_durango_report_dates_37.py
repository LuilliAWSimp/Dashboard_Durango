from io import BytesIO
from unittest.mock import patch

from openpyxl import load_workbook

from app.services.water_daily_report_service import build_daily_water_report_excel, build_daily_water_report_pdf, get_daily_water_report


def item(name: str, module: str, key: str, volume: float) -> dict:
    return {
        'name': name, 'module': module, 'operational_key': key,
        'current_flow': 1.0, 'flow_unit': 'L/s', 'period_open_m3': 10.0,
        'period_close_m3': 10.0 + volume, 'period_m3': volume,
        'period_m3_reliable': True, 'validated_volume_m3': volume,
        'activity': 'Con actividad', 'communication': 'Actualizado',
        'last_update': '2026-09-25T12:00:00', 'samples': 10,
    }


def comparison(date_value: str, wells: float, lines: float, washers: float, jarabes: float) -> dict:
    return {
        'date': date_value,
        'operational_groups': {
            'wells': {'subtotal_validated_m3': wells},
            'lines': {'subtotal_validated_m3': lines},
            'lavadoras': {'subtotal_validated_m3': washers},
            'jarabes': {'subtotal_validated_m3': jarabes},
        },
    }


def review_payload() -> dict:
    wells = [item('Pozo 1', 'well', 'pozo_1', 10.0)]
    lines = [item('Línea 1', 'line', 'linea_1', 8.0)]
    flows = [item('Lavadora Línea 2', 'flow', 'lavadora_linea_2', 4.0), item('Jarabes', 'flow', 'jarabes', 2.0)]
    return {
        'source_status': 'operational',
        'modules': {
            'wells': {'items': wells, 'summary': {'active_count': 1}},
            'lines': {'items': lines, 'summary': {'active_count': 1}},
            'flows': {'items': flows, 'summary': {'active_count': 2}},
        },
        'comparatives': {
            'previous_day': comparison('2026-09-24', 9.0, 7.0, 3.0, 1.0),
            'previous_week': comparison('2026-09-18', 8.0, 6.0, 2.0, 0.5),
        },
        'shifts': {'shifts': []},
    }


def build_report() -> dict:
    with patch('app.services.water_daily_report_service.get_daily_water_review', return_value=review_payload()), patch(
        'app.services.water_daily_report_service.get_water_history_module', return_value={'series': [], 'aggregation': 'quarter_hour'}
    ):
        return get_daily_water_report('2026-09-25')


def test_daily_report_exposes_explicit_comparison_dates() -> None:
    report = build_report()
    headers = report['comparatives']['headers']
    assert headers == {
        'selected': 'Seleccionado · 25/09/2026',
        'previous': 'Anterior · 24/09/2026',
        'previous_week': 'Semana anterior · 18/09/2026',
    }
    rows = {row['key']: row for row in report['comparatives']['rows']}
    assert rows['wells']['selected_m3'] == 10.0
    assert rows['wells']['previous_m3'] == 9.0
    assert rows['wells']['previous_week_m3'] == 8.0


def test_excel_uses_explicit_period_headers_and_comparison() -> None:
    report = build_report()
    report['generated_at'] = '2026-09-25T15:45:30'
    content, filename = build_daily_water_report_excel(report)
    workbook = load_workbook(BytesIO(content), data_only=False)
    assert workbook['Pozos']['A1'].value == 'Pozo'
    assert workbook['Pozos']['C1'].value == 'Totalizador apertura (m³)'
    assert workbook['Pozos']['D1'].value == 'Totalizador al cierre (m³)'
    assert workbook['Pozos']['E1'].value == 'Volumen bombeado · 25/09/2026 (m³)'
    values = [workbook['Resumen'].cell(row, 1).value for row in range(1, workbook['Resumen'].max_row + 1)]
    assert 'Comparativo de volumen por módulo' in values
    assert filename == 'reporte-control-hidrico-durango-2026-09-25_generado-15-45-30.xlsx'


def test_pdf_filename_uses_period_and_generation_time() -> None:
    report = build_report()
    report['generated_at'] = '2026-09-25T15:45:30'
    content, filename = build_daily_water_report_pdf(report)
    assert content.startswith(b'%PDF')
    assert filename == 'reporte-control-hidrico-durango-2026-09-25_generado-15-45-30.pdf'


def test_range_contract_exposes_full_bounds() -> None:
    from datetime import date
    from app.services.water_daily_report_service import _comparison_period_contract, _report_period_metric_header
    contract = _comparison_period_contract(date(2026, 9, 19), date(2026, 9, 25))
    assert contract['headers']['selected'] == 'Seleccionado · 19/09/2026 → 25/09/2026'
    assert contract['headers']['previous'] == 'Anterior · 18/09/2026 → 24/09/2026'
    assert contract['headers']['previous_week'] == 'Semana anterior · 12/09/2026 → 18/09/2026'
    report = {'start_date': '2026-09-19', 'end_date': '2026-09-25'}
    assert _report_period_metric_header(report, 'Volumen consumido') == 'Volumen consumido · 19/09/2026 → 25/09/2026'


def test_range_filename_contains_both_limits() -> None:
    from app.services.water_daily_report_service import _report_filename_stem
    report = {'start_date': '2026-09-19', 'end_date': '2026-09-25', 'generated_at': '2026-09-25T15:45:30'}
    assert _report_filename_stem(report) == 'reporte-control-hidrico-durango-2026-09-19_a_2026-09-25_generado-15-45-30'
