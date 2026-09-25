from io import BytesIO

from openpyxl import load_workbook

import app.services.water_five_minute_export_service as service


def _payload(name: str, element_id, operational_key: str, volume: float):
    return {
        'plant': 'Planta Durango',
        'module': 'flow',
        'element_id': element_id,
        'sensor_id': element_id if isinstance(element_id, int) else None,
        'operational_key': operational_key,
        'name': name,
        'flow_unit': 'L/s',
        'start_date': '2026-09-20',
        'end_date': '2026-09-20',
        'effective_start_local': '2026-09-20T00:00:00',
        'effective_end_local': '2026-09-21T00:00:00',
        'time_zone': 'America/Monterrey',
        'source_status': 'prueba',
        'rows': [
            {
                'element': name,
                'element_id': element_id,
                'start_local': None,
                'end_local': None,
                'flow_avg': 1.25,
                'flow_min': 1.0,
                'flow_max': 1.5,
                'totalizer_open_m3': 10.0,
                'totalizer_close_m3': 10.0 + volume,
                'validated_volume_m3': volume,
                'reported_volume_m3': volume,
                'volume_reliable': True,
                'samples': 5,
                'samples_expected': 5,
                'coverage_pct': 100.0,
                'quality_label': 'Validado',
                'opening_source': 'previous',
            }
        ],
    }


def test_build_module_five_minute_excel_creates_summary_and_one_sheet_per_element():
    content, filename = service.build_five_minute_module_excel(
        module='flow',
        payloads=[
            _payload('Lavadora Línea 2', 2004, 'lavadora_linea_2', 1.2),
            _payload('Lavadora Vidrio', 'lavadora_vidrio', 'lavadora_vidrio', 0.8),
        ],
        start_date='2026-09-20',
        end_date='2026-09-20',
        view_label='Lavadoras',
    )
    workbook = load_workbook(BytesIO(content), data_only=True)
    assert workbook.sheetnames[0] == 'Resumen'
    assert 'Lavadora Línea 2' in workbook.sheetnames
    assert 'Lavadora Vidrio' in workbook.sheetnames
    assert workbook['Resumen']['A1'].value == 'ARCA Durango - Lavadoras - Excel 5 minutos'
    assert workbook['Resumen']['G2'].value == 2
    assert filename == 'ARCA_Durango_Lavadoras_5min_2026-09-20_2026-09-20.xlsx'


def test_export_module_five_minute_accepts_numeric_and_operational_key_and_deduplicates(monkeypatch):
    calls = []

    def fake_payload(*, module, element_id, start_date, end_date):
        calls.append((module, str(element_id)))
        name = 'Lavadora Línea 2' if str(element_id) == '2004' else 'Lavadora Vidrio'
        return _payload(name, int(element_id) if str(element_id).isdigit() else str(element_id), str(element_id), 1.0)

    monkeypatch.setattr(service, 'get_five_minute_export_data', fake_payload)
    content, filename = service.export_five_minute_module_excel(
        module='flow',
        element_ids=['2004', 'lavadora_vidrio', '2004'],
        start_date='2026-09-20',
        end_date='2026-09-20',
        view_label='Lavadoras',
    )

    assert calls == [('flow', '2004'), ('flow', 'lavadora_vidrio')]
    assert filename.startswith('ARCA_Durango_Lavadoras_5min_')
    workbook = load_workbook(BytesIO(content), data_only=True)
    assert workbook['Resumen']['G2'].value == 2


def test_module_export_rejects_empty_selection():
    try:
        service.export_five_minute_module_excel(
            module='flow',
            element_ids=[],
            start_date='2026-09-20',
            end_date='2026-09-20',
            view_label='Lavadoras',
        )
    except ValueError as exc:
        assert 'Selecciona al menos un elemento' in str(exc)
    else:
        raise AssertionError('La selección vacía debía rechazarse.')
